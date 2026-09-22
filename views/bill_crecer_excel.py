import re
import unicodedata
from datetime import date, datetime
from typing import Any, Dict, List

from openpyxl import load_workbook

# Columnas esperadas para el formato actual de Crecer
ESTADO_CUENTA_COLUMNS = (
    "USUARIO",
    "CORRELATIVO",
    "PRODUCTO",
    "MOVIMIENTO",
    "PÓLIZA",
    "CONTRATANTE",
    "FECHA DE REGISTRO",
    "PRIMA TOTAL",
    "MEDIO DE PAGO",
    "NRO DE COMPROBANTE",
    "CÓDIGO DE PAGO",
    "ESTADO DE PAGO",
    "FECHA DE PAGO",
)

_HEADER_HINTS = {
    "USUARIO": ("USUARIO", "USUARI"),
    "CORRELATIVO": ("CORRELATIVO",),
    "PRODUCTO": ("PRODUCTO",),
    "MOVIMIENTO": ("MOVIMIENTO",),
    "POLIZA": ("POLIZA",),
    "CONTRATANTE": ("CONTRATANTE",),
    "FECHA DE REGISTRO": ("FECHA DE REGISTRO", "FECHA REGISTRO"),
    "PRIMA TOTAL": ("PRIMA TOTAL",),
    "MEDIO DE PAGO": ("MEDIO DE PAGO",),
    "NRO DE COMPROBANTE": (
        "NRO DE COMPROBANTE",
        "NRO COMPROBANTE",
    ),
    "CODIGO DE PAGO": (
        "CODIGO DE PAGO",
        "CODIGO PAGO",
    ),
    "ESTADO DE PAGO": (
        "ESTADO DE PAGO",
        "ESTADO PAGO",
    ),
    "FECHA DE PAGO": ("FECHA DE PAGO", "FECHA PAGO"),
}

_HEADER_SCAN_MAX_ROWS = 60


def _normalizar_header(valor: Any) -> str:
    if valor is None:
        return ""

    texto = str(valor).strip().upper()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))

    texto = texto.replace("°", "")
    texto = texto.replace(".", " ")
    texto = texto.replace("_", " ")

    return re.sub(r"\s+", " ", texto).strip()


def _limpiar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, str):
        return re.sub(r"\s+", " ", valor.strip())
    return str(valor).strip()


def _limpiar_numero(valor: Any) -> float:
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()
    if not texto:
        return 0.0

    texto = (
        texto.replace("S/.", "")
        .replace("S/", "")
        .replace(",", "")
        .strip()
    )

    m = re.search(r"-?\d+(?:\.\d+)?", texto)
    if not m:
        return 0.0

    try:
        return float(m.group(0))
    except Exception:
        return 0.0


def _formatear_fecha(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")

    texto = str(valor).strip()
    if not texto:
        return ""

    m = re.match(r"^(\d{1,2})[\/\-\.](\d{1,2})[\/\-\.](\d{4})$", texto)
    if m:
        d, mes, anio = m.groups()
        return f"{int(d):02d}/{int(mes):02d}/{anio}"

    m = re.match(r"^(\d{4})[\/\-\.](\d{1,2})[\/\-\.](\d{1,2})$", texto)
    if m:
        anio, mes, d = m.groups()
        return f"{int(d):02d}/{int(mes):02d}/{anio}"

    return texto


def _buscar_columna(mapa: Dict[str, int], *nombres: str):
    for nombre in nombres:
        clave = _normalizar_header(nombre)
        if clave in mapa:
            return mapa[clave]
    return None


def _match_header(header: str, *hints: str) -> bool:
    texto = _normalizar_header(header)
    if not texto:
        return False
    return any(
        _normalizar_header(hint) in texto
        or texto in _normalizar_header(hint)
        for hint in hints
    )


def _detectar_headers(row: Any) -> Dict[str, int]:
    detected: Dict[str, int] = {}
    for idx, cell in enumerate(row):
        header = _normalizar_header(cell)
        if not header:
            continue
        for canonical, hints in _HEADER_HINTS.items():
            if canonical in detected:
                continue
            if _match_header(header, *hints):
                detected[canonical] = idx
                break
    return detected


def _es_fila_total_resumen(
    row: Any,
    correlativo: str,
    poliza: str,
    contratante: str,
    nro_comprobante: str,
    codigo_pago: str,
) -> bool:
    valores = [
        _normalizar_header(cell)
        for cell in row
        if cell is not None and str(cell).strip()
    ]
    if not valores:
        return False

    primer_valor = valores[0]
    es_marcador_total = (
        primer_valor == "TOTAL"
        or primer_valor.startswith("TOTAL ")
        or primer_valor == "SUBTOTAL"
        or primer_valor.startswith("SUBTOTAL ")
        or primer_valor == "TOTAL GENERAL"
    )
    if not es_marcador_total:
        return False

    # Si hay identificadores reales de negocio, es una fila válida aunque
    # una empresa tenga la palabra TOTAL en su nombre.
    if any([correlativo, poliza, contratante, nro_comprobante, codigo_pago]):
        return False

    return True


def _puntaje_headers(detected: Dict[str, int], row_number: int) -> tuple:
    return (
        len(detected),
        1 if "USUARIO" in detected else 0,
        row_number,
    )


def extraer_estado_cuenta_crecer(excel_path: str) -> List[Dict[str, Any]]:
    workbook = load_workbook(
        filename=excel_path,
        data_only=True,
        read_only=True,
    )

    try:
        mejor_detectado: Dict[str, int] = {}
        resultado = []
        for worksheet in workbook.worksheets:
            header_row_number = None
            mapa = {}
            mejor_detectado_hoja: Dict[str, int] = {}
            mejor_fila_hoja = None
            mejor_puntaje_hoja = (-1, -1, -1)

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(
                        worksheet.max_row or _HEADER_SCAN_MAX_ROWS,
                        _HEADER_SCAN_MAX_ROWS,
                    ),
                    values_only=True,
                ),
                start=1,
            ):
                detected = _detectar_headers(row)
                if len(detected) > len(mejor_detectado):
                    mejor_detectado = detected
                puntaje = _puntaje_headers(detected, row_number)
                if puntaje > mejor_puntaje_hoja:
                    mejor_detectado_hoja = detected
                    mejor_fila_hoja = row_number
                    mejor_puntaje_hoja = puntaje

            if len(mejor_detectado_hoja) >= 4:
                header_row_number = mejor_fila_hoja
                mapa = mejor_detectado_hoja

            if header_row_number is None:
                continue

            col_correlativo = mapa.get("CORRELATIVO")
            col_usuario = mapa.get("USUARIO")
            col_producto = mapa.get("PRODUCTO")
            col_movimiento = mapa.get("MOVIMIENTO")
            col_poliza = mapa.get("POLIZA")
            col_contratante = mapa.get("CONTRATANTE")
            col_fecha_registro = mapa.get("FECHA DE REGISTRO")
            col_prima_total = mapa.get("PRIMA TOTAL")
            col_medio_pago = mapa.get("MEDIO DE PAGO")
            col_nro_comprobante = mapa.get("NRO DE COMPROBANTE")
            col_codigo_pago = mapa.get("CODIGO DE PAGO")
            col_estado_pago = mapa.get("ESTADO DE PAGO")
            col_fecha_pago = mapa.get("FECHA DE PAGO")

            for row in worksheet.iter_rows(
                min_row=header_row_number + 1,
                values_only=True,
            ):
                if not row:
                    continue

                def obtener(col):
                    if col is None or col >= len(row):
                        return ""
                    return row[col]

                correlativo = _limpiar_texto(obtener(col_correlativo))
                usuario = _limpiar_texto(obtener(col_usuario))
                producto = _limpiar_texto(obtener(col_producto))
                movimiento = _limpiar_texto(obtener(col_movimiento))
                poliza = _limpiar_texto(obtener(col_poliza))
                contratante = _limpiar_texto(obtener(col_contratante))
                fecha_registro = _formatear_fecha(obtener(col_fecha_registro))
                prima_total = _limpiar_numero(obtener(col_prima_total))
                medio_pago = _limpiar_texto(obtener(col_medio_pago))
                nro_comprobante = _limpiar_texto(obtener(col_nro_comprobante))
                codigo_pago = _limpiar_texto(obtener(col_codigo_pago))
                estado_pago = _limpiar_texto(obtener(col_estado_pago))
                fecha_pago = _formatear_fecha(obtener(col_fecha_pago))

                if not any([
                    usuario, correlativo, producto, movimiento, poliza, contratante,
                    fecha_registro, prima_total, medio_pago, nro_comprobante,
                    codigo_pago, estado_pago, fecha_pago
                ]):
                    continue

                if _es_fila_total_resumen(
                    row,
                    correlativo,
                    poliza,
                    contratante,
                    nro_comprobante,
                    codigo_pago,
                ):
                    continue

                resultado.append({
                    "usuario": usuario,
                    "correlativo": correlativo,
                    "producto": producto,
                    "movimiento": movimiento,
                    "poliza": poliza,
                    "contratante": contratante,
                    "fecha_registro": fecha_registro,
                    "prima_total": prima_total,
                    "medio_pago": medio_pago,
                    "nro_comprobante": nro_comprobante,
                    "codigo_pago": codigo_pago,
                    "estado_pago": estado_pago,
                    "fecha_pago": fecha_pago,
                })

        if resultado:
            return resultado

        raise ValueError(
            "No se encontró una hoja con los encabezados válidos para el Estado de Cuenta de Crecer.\n"
            f"Encabezados detectados: {', '.join(mejor_detectado.keys()) or 'ninguno'}"
        )

    finally:
        workbook.close()


def extraer_estado_cuenta_sanitas(excel_path: str) -> List[Dict[str, Any]]:
    # Compatibilidad temporal con llamadas antiguas.
    return extraer_estado_cuenta_crecer(excel_path)
